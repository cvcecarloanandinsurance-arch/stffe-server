import io, os, gc, re, threading
from flask import Flask, request, jsonify, send_file, make_response
from werkzeug.exceptions import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

VERSION = "1.1"
API_TOKEN = os.environ.get("API_TOKEN", "").strip()
MAX_SIDE = int(os.environ.get("OCR_MAX_SIDE", "2000"))      # huge photos are shrunk -> far less RAM
OCR_TIMEOUT = int(os.environ.get("OCR_TIMEOUT", "100"))     # seconds; a stuck Tesseract is killed
BUSY_WAIT = int(os.environ.get("BUSY_WAIT", "90"))          # wait this long for the single work slot
app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024
heavy = threading.BoundedSemaphore(1)   # ONE OCR/Excel job at a time -> RAM never stacks up

try:
    import pytesseract
    from PIL import Image, ImageOps
    Image.MAX_IMAGE_PIXELS = 50_000_000
    LANGS = [l for l in pytesseract.get_languages(config="") if l != "osd"]
    OCR_OK = True
except Exception:
    LANGS, OCR_OK = [], False


@app.after_request
def cors(r):
    r.headers["Access-Control-Allow-Origin"] = "*"
    r.headers["Access-Control-Allow-Headers"] = "X-Token, Content-Type"
    r.headers["Access-Control-Allow-Methods"] = "GET, POST, OPTIONS"
    r.headers["Access-Control-Max-Age"] = "86400"
    return r


@app.errorhandler(Exception)
def on_error(e):   # always JSON (+CORS) so the app shows the real reason instead of "connection lost"
    if isinstance(e, HTTPException):
        return jsonify(ok=False, error=e.description), e.code
    return jsonify(ok=False, error=str(e)[:200]), 500


@app.route("/<path:p>", methods=["OPTIONS"])
@app.route("/", methods=["OPTIONS"])
def opts(p=""):
    return make_response("", 204)


def authed():
    return bool(API_TOKEN) and request.headers.get("X-Token", "").strip() == API_TOKEN


@app.route("/")
def home():
    return "STFFE server running"


@app.route("/health")
def health():
    mem = None
    try:
        import resource
        mem = round(resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024)
    except Exception:
        pass
    return jsonify(ok=True, auth=authed(), version=VERSION, ocr=OCR_OK, langs=LANGS, mem_mb=mem)


def run_ocr(img, lang, cfg):
    # ONE Tesseract run (old code ran it twice: text + data) -> half the time and memory
    d = pytesseract.image_to_data(img, lang=lang, config=cfg, timeout=OCR_TIMEOUT,
                                 output_type=pytesseract.Output.DICT)
    lines, cur, confs, last_line, last_par = [], [], [], None, None
    for i, t in enumerate(d["text"]):
        if not str(t).strip():
            continue
        try:
            c = float(d["conf"][i])
            if c >= 0:
                confs.append(c)
        except Exception:
            pass
        par = (d["block_num"][i], d["par_num"][i])
        line = par + (d["line_num"][i],)
        if line != last_line:
            if cur:
                lines.append(" ".join(cur))
            cur = []
            if last_par is not None and par != last_par:
                lines.append("")
            last_line, last_par = line, par
        cur.append(str(t).strip())
    if cur:
        lines.append(" ".join(cur))
    return "\n".join(lines) + "\n", (sum(confs) / len(confs) if confs else 0)


@app.route("/ocr", methods=["POST"])
def ocr():
    if not authed():
        return jsonify(ok=False, error="unauthorized"), 401
    if not OCR_OK:
        return jsonify(ok=False, error="tesseract missing"), 500
    f = request.files.get("file")
    if not f:
        return jsonify(ok=False, error="no file"), 400
    if not heavy.acquire(timeout=BUSY_WAIT):
        return jsonify(ok=False, error="server busy, retry"), 503
    img = None
    try:
        im = Image.open(f.stream)
        try:
            im.draft("RGB", (MAX_SIDE, MAX_SIDE))      # JPEG: decode small from the start
        except Exception:
            pass
        img = ImageOps.exif_transpose(im).convert("RGB")
        if max(img.size) > MAX_SIDE:
            img.thumbnail((MAX_SIDE, MAX_SIDE))
        lang = "+".join([l for l in ("eng", "tam") if l in LANGS]) or "eng"
        psm = request.args.get("psm", "").strip()
        cfg = "--oem 1" + (" --psm " + psm if psm.isdigit() and int(psm) <= 13 else "")
        text, conf = run_ocr(img, lang, cfg)
        return jsonify(ok=True, text=text, confidence=conf)
    except Image.DecompressionBombError:
        return jsonify(ok=False, error="image too large"), 413
    except RuntimeError as e:
        if "timeout" in str(e).lower():
            return jsonify(ok=False, error="OCR timeout"), 504
        return jsonify(ok=False, error=str(e)[:200]), 500
    except Exception as e:
        return jsonify(ok=False, error=str(e)[:200]), 500
    finally:
        img = None
        gc.collect()
        heavy.release()


def sheet_name(n, used):
    n = re.sub(r"[\[\]\*\?/\\:]", " ", str(n or "Sheet")).strip()[:31] or "Sheet"
    base, k = n, 2
    while n.lower() in used:
        n = base[:28] + " " + str(k)
        k += 1
    used.add(n.lower())
    return n


def cell_val(v):
    if v is None:
        return ""
    return str(v) if isinstance(v, (list, dict, tuple, set)) else v


@app.route("/report/xlsx", methods=["POST"])
def xlsx():
    if not authed():
        return jsonify(ok=False, error="unauthorized"), 401
    spec = request.get_json(force=True, silent=True) or {}
    if not heavy.acquire(timeout=BUSY_WAIT):
        return jsonify(ok=False, error="server busy, retry"), 503
    try:
        wb = Workbook()
        wb.remove(wb.active)
        used = set()
        for sh in spec.get("sheets", []) or [{"name": "Sheet1", "columns": [], "rows": []}]:
            ws = wb.create_sheet(sheet_name(sh.get("name"), used))
            cols = sh.get("columns", []) or []
            ws.append([c.get("header", "") for c in cols])
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            widths = [len(str(c.get("header", ""))) for c in cols]
            for row in sh.get("rows", []) or []:
                if not isinstance(row, dict):
                    continue
                vals = [cell_val(row.get(c.get("key"), "")) for c in cols]
                ws.append(vals)
                for i, v in enumerate(vals):
                    if len(str(v)) > widths[i]:
                        widths[i] = len(str(v))
            for i, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 10), 45)
            for r in ws.iter_rows(min_row=2):
                for cell in r:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.freeze_panes = "A2"
            if sh.get("landscape"):
                ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        name = re.sub(r"[^\w.\- ]", "_", str(spec.get("filename") or "report.xlsx"))
        return send_file(buf, as_attachment=True, download_name=name,
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    finally:
        gc.collect()
        heavy.release()
